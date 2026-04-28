from __future__ import annotations

from pathlib import Path
from typing import Optional, Dict, Any, List
import mimetypes
import subprocess
import sys

from schemas.evidence_models import RawEvidence
from tools.file_fingerprint import sha256_file
from tools.pdf_extractor import PdfExtractor

_pdf_extractor = PdfExtractor()


TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".yaml", ".yml"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"}

# 新增：Excel 文件支持
EXCEL_X_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
EXCEL_LEGACY_EXTENSIONS = {".xls"}


class ParseError(Exception):
    pass


def _read_text_file(path: Path) -> str:
    encodings = ["utf-8", "utf-8-sig", "gb18030", "latin-1"]
    for enc in encodings:
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ParseError(f"无法读取文本编码: {path}")


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)

    return str(value).strip()


def _trim_empty_tail(values: List[str]) -> List[str]:
    values = list(values)

    while values and not str(values[-1]).strip():
        values.pop()

    return values


def _read_xlsx(path: Path, max_rows_per_sheet: int = 2000, max_cols: int = 80) -> str:
    """
    读取 .xlsx / .xlsm / .xltx / .xltm

    输出格式：
    [Sheet] Sheet1
    A列\tB列\tC列...
    ...
    """
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise ParseError("读取 XLSX/XLSM 需要安装 openpyxl：pip install openpyxl") from e

    try:
        wb = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
        )
    except Exception as e:
        raise ParseError(f"XLSX/XLSM 解析失败: {e}") from e

    parts: List[str] = []

    try:
        for ws in wb.worksheets:
            parts.append(f"[Sheet] {ws.title}")

            row_count = 0

            for row in ws.iter_rows(values_only=True):
                row_count += 1

                if row_count > max_rows_per_sheet:
                    parts.append(f"... 已截断：该 Sheet 超过 {max_rows_per_sheet} 行")
                    break

                cells = [_cell_to_text(v) for v in list(row)[:max_cols]]
                cells = _trim_empty_tail(cells)

                if not cells:
                    continue

                parts.append("\t".join(cells))

            parts.append("")

    finally:
        try:
            wb.close()
        except Exception:
            pass

    return "\n".join(parts).strip()


def _read_xls(path: Path, max_rows_per_sheet: int = 2000, max_cols: int = 80) -> str:
    """
    读取老版本 .xls。

    注意：
    xlrd 2.x 只支持 .xls，不支持 .xlsx。
    安装：
    pip install xlrd
    """
    try:
        import xlrd
    except Exception as e:
        raise ParseError("读取 XLS 需要安装 xlrd：pip install xlrd") from e

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as e:
        raise ParseError(f"XLS 解析失败: {e}") from e

    parts: List[str] = []

    for sheet in book.sheets():
        parts.append(f"[Sheet] {sheet.name}")

        row_limit = min(sheet.nrows, max_rows_per_sheet)

        for r in range(row_limit):
            cells = []

            col_limit = min(sheet.ncols, max_cols)

            for c in range(col_limit):
                cell = sheet.cell(r, c)
                value = cell.value

                # 日期类型简单保留原始值；复杂格式后面需要再专项优化
                cells.append(_cell_to_text(value))

            cells = _trim_empty_tail(cells)

            if not cells:
                continue

            parts.append("\t".join(cells))

        if sheet.nrows > max_rows_per_sheet:
            parts.append(f"... 已截断：该 Sheet 超过 {max_rows_per_sheet} 行")

        parts.append("")

    return "\n".join(parts).strip()


def _read_pdf(path: Path) -> Dict[str, Any]:
    """
    PDF 统一读取：
    1. 先尝试原生文本抽取
    2. 文本不足时自动切 OCR
    """
    try:
        result = _pdf_extractor.extract(str(path))
        return {
            "ok": result.get("ok", False),
            "text": result.get("text", ""),
            "method": result.get("method", "text_pdf"),
            "page_count": result.get("page_count", 0),
            "warnings": result.get("warnings", []),
        }
    except Exception as e:
        raise ParseError(f"PDF 解析失败: {e}") from e


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception as e:
        raise ParseError("读取 DOCX 需要安装 python-docx") from e

    try:
        doc = Document(str(path))
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        raise ParseError(f"DOCX 解析失败: {e}") from e


def _read_doc_by_word(path: Path) -> str:
    """
    Windows 下优先使用 Word COM 读取 .doc
    需要：
    - 安装 pywin32
    - 本机安装 Microsoft Word
    """
    if not sys.platform.startswith("win"):
        raise ParseError("当前不是 Windows，无法使用 Word COM 读取 DOC")

    try:
        import win32com.client  # type: ignore
    except Exception as e:
        raise ParseError("读取 DOC 需要安装 pywin32") from e

    word = None
    doc = None

    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0

        doc = word.Documents.Open(str(path.resolve()))
        text = doc.Content.Text
        return text or ""

    except Exception as e:
        raise ParseError(f"Word COM 解析 DOC 失败: {e}") from e

    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass

        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass


def _read_doc_fallback(path: Path) -> str:
    for cmd in (["antiword", str(path)], ["catdoc", str(path)]):
        try:
            result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
            if result.returncode == 0:
                return result.stdout
        except FileNotFoundError:
            continue
        except Exception:
            continue

    raise ParseError("DOC fallback 解析失败：未找到 antiword/catdoc，或执行失败")


def _read_doc(path: Path) -> str:
    # 优先 Word COM，其次 antiword/catdoc
    try:
        return _read_doc_by_word(path)
    except ParseError:
        return _read_doc_fallback(path)


def _ocr_image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except Exception as e:
        raise ParseError("图片 OCR 需要安装 pillow + pytesseract") from e

    try:
        img = Image.open(path)
        return pytesseract.image_to_string(img)
    except Exception as e:
        raise ParseError(f"图片 OCR 失败: {e}") from e


def _looks_like_noise(text: str) -> bool:
    cleaned = text.strip()

    if not cleaned:
        return True

    if len(cleaned) < 6:
        return True

    printable = sum(ch.isprintable() for ch in cleaned)
    ratio_printable = printable / max(len(cleaned), 1)

    alnum_or_cjk = sum(ch.isalnum() or ("\u4e00" <= ch <= "\u9fff") for ch in cleaned)
    ratio_signal = alnum_or_cjk / max(len(cleaned), 1)

    return ratio_printable < 0.7 or ratio_signal < 0.15


def parse_document(file_path: str, try_ocr: bool = False) -> RawEvidence:
    path = Path(file_path)

    if not path.exists():
        raise ParseError(f"文件不存在: {file_path}")

    if not path.is_file():
        raise ParseError(f"不是文件: {file_path}")

    ext = path.suffix.lower()
    mime_type, _ = mimetypes.guess_type(str(path))
    size = path.stat().st_size
    sha256 = sha256_file(path)

    extracted_text: Optional[str] = None
    is_binary = True
    metadata: Dict[str, Any] = {"try_ocr": try_ocr}

    if size == 0:
        return RawEvidence(
            file_path=str(path),
            file_name=path.name,
            extension=ext,
            mime_type=mime_type,
            size=size,
            sha256=sha256,
            extracted_text="",
            is_binary=False,
            metadata=metadata,
        )

    try:
        if ext in TEXT_EXTENSIONS:
            extracted_text = _read_text_file(path)
            is_binary = False

        elif ext == ".pdf":
            pdf_result = _read_pdf(path)
            extracted_text = pdf_result["text"]

            metadata["pdf_extraction_method"] = pdf_result.get("method")
            metadata["pdf_page_count"] = pdf_result.get("page_count", 0)

            if pdf_result.get("warnings"):
                metadata["pdf_warnings"] = pdf_result["warnings"]

            if not pdf_result.get("ok", False) and not (extracted_text or "").strip():
                metadata["parser_warning"] = "PDF 经原生解析/OCR 后仍无可用内容"
                metadata["text_extract_failed"] = True

        elif ext == ".docx":
            extracted_text = _read_docx(path)

        elif ext == ".doc":
            try:
                extracted_text = _read_doc(path)
            except ParseError as e:
                extracted_text = None
                metadata["parser_warning"] = str(e)
                metadata["text_extract_failed"] = True

        elif ext in EXCEL_X_EXTENSIONS:
            try:
                extracted_text = _read_xlsx(path)
                metadata["excel_extraction_method"] = "openpyxl"
            except ParseError as e:
                extracted_text = None
                metadata["parser_warning"] = str(e)
                metadata["text_extract_failed"] = True

        elif ext in EXCEL_LEGACY_EXTENSIONS:
            try:
                extracted_text = _read_xls(path)
                metadata["excel_extraction_method"] = "xlrd"
            except ParseError as e:
                extracted_text = None
                metadata["parser_warning"] = str(e)
                metadata["text_extract_failed"] = True

        elif ext in IMAGE_EXTENSIONS:
            if try_ocr:
                try:
                    extracted_text = _ocr_image(path)
                except ParseError as e:
                    extracted_text = None
                    metadata["parser_warning"] = str(e)
                    metadata["text_extract_failed"] = True
            else:
                extracted_text = None
                metadata["ocr_skipped"] = True

        else:
            extracted_text = None
            metadata["unsupported_text_extraction"] = True

    except ParseError:
        raise

    except Exception as e:
        extracted_text = None
        metadata["parser_warning"] = f"未知解析异常: {e}"
        metadata["text_extract_failed"] = True

    return RawEvidence(
        file_path=str(path),
        file_name=path.name,
        extension=ext,
        mime_type=mime_type,
        size=size,
        sha256=sha256,
        extracted_text=extracted_text,
        is_binary=is_binary,
        metadata=metadata,
    )


def is_empty_evidence(raw: RawEvidence) -> bool:
    if raw.size == 0:
        return True

    if raw.extracted_text is None:
        return False

    return len(raw.extracted_text.strip()) == 0


def is_noise_evidence(raw: RawEvidence) -> bool:
    if raw.extracted_text is None:
        return False

    return _looks_like_noise(raw.extracted_text)